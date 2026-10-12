"""Inject computed formula values into the workbook's cached values.

Reads the existing xlsx, computes every formula using the `formulas` library,
then rewrites each sheet XML so each formula cell carries both the formula
and the cached numeric/string result.
"""

from __future__ import annotations

import pathlib
import re
import zipfile

import formulas
from lxml import etree


HERE = pathlib.Path(__file__).parent
SRC = HERE / "test_data" / "complex_workbook.xlsx"
DST = HERE / "test_data" / "complex_workbook_cached.xlsx"


NS = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHEET_KEY_RE = re.compile(r"^'\[([^\]]+)\]([^!]+)!([^!]+)$")


def compute_values() -> dict[str, dict[str, object]]:
    xl = formulas.ExcelModel().loads(str(SRC)).finish()
    sol = xl.calculate()
    out: dict[str, dict[str, object]] = {}
    for k, node in sol.items():
        m = SHEET_KEY_RE.match(k)
        if not m:
            continue
        sheet = m.group(2).rstrip("'")
        coord = m.group(3)
        try:
            val = node.value
            if hasattr(val, "tolist"):
                val = val.tolist()
            if isinstance(val, list) and len(val) == 1 and isinstance(val[0], list) and len(val[0]) == 1:
                val = val[0][0]
        except Exception:
            val = None
        out.setdefault(sheet, {})[coord] = val
    return out


def _serialise_value(v):
    if v is None:
        return ("str", "")
    if isinstance(v, bool):
        return ("b", "1" if v else "0")
    if isinstance(v, (int, float)):
        return ("n", repr(float(v)))
    import datetime as dt
    if isinstance(v, (dt.date, dt.datetime)):
        base = dt.datetime(1899, 12, 30)
        serial = (v - base).total_seconds() / 86400.0 if isinstance(v, dt.datetime) else (v - dt.date(1899, 12, 30)).days
        return ("n", repr(float(serial)))
    return ("str", str(v))


def inject_into_xml(sheet_xml: bytes, computed: dict[str, object]) -> bytes:
    root = etree.fromstring(sheet_xml)
    ns_main = NS["main"]
    for c in root.iter(f"{{{ns_main}}}c"):
        coord = c.get("r")
        f = c.find(f"{{{ns_main}}}f")
        if f is None or coord is None:
            continue
        for v in c.findall(f"{{{ns_main}}}v"):
            c.remove(v)
        computed_val = computed.get(coord)
        sv = _serialise_value(computed_val)
        if sv is None:
            continue
        t_attr, val_text = sv
        if t_attr == "b":
            c.set("t", "b")
        elif t_attr == "str":
            c.set("t", "str")
        else:
            if c.get("t") in ("str", "b", "s"):
                del c.attrib["t"]
        new_v = etree.SubElement(c, f"{{{ns_main}}}v")
        new_v.text = val_text
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def main():
    computed = compute_values()
    sheet_name_to_target: dict[str, str] = {}

    with zipfile.ZipFile(SRC, "r") as zin:
        wb_xml = zin.read("xl/workbook.xml")
        wb_rels = zin.read("xl/_rels/workbook.xml.rels")

    wb_root = etree.fromstring(wb_xml)
    rels_root = etree.fromstring(wb_rels)
    rid_to_target = {r.get("Id"): r.get("Target") for r in rels_root}

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    for sh in wb_root.iter(f"{{{ns}}}sheet"):
        name = sh.get("name")
        rid = sh.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rid_to_target[rid]
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("xl/"):
            target = "xl/" + target
        sheet_name_to_target[name] = target

    print("Sheet -> file mapping:")
    for k, v in sheet_name_to_target.items():
        print(f"  {k!r} -> {v!r}")

    # Build case-insensitive lookup: formulas library uppercases sheet names
    computed_ci: dict[str, dict[str, object]] = {k.upper(): v for k, v in computed.items()}

    DST.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(SRC, "r") as zin, zipfile.ZipFile(DST, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            sheet_match = None
            for name, target in sheet_name_to_target.items():
                if item == target:
                    sheet_match = name
                    break
            if sheet_match is not None and item.endswith(".xml"):
                sheet_computed = computed_ci.get(sheet_match.upper(), {})
                new_data = inject_into_xml(data, sheet_computed)
                zout.writestr(item, new_data)
            else:
                zout.writestr(item, data)

    print(f"Wrote {DST}")

    import openpyxl
    wb2 = openpyxl.load_workbook(DST, data_only=True)
    print("Verification (data_only=True):")
    print("  Sales F2  :", wb2["Sales"]["F2"].value, " (expected 62.5)")
    print("  Sales F42 :", wb2["Sales"]["F42"].value)
    print("  Sales H2  :", wb2["Sales"]["H2"].value)
    print("  Dashboard B4:", wb2["Dashboard"]["B4"].value)
    print("  Dashboard B6:", wb2["Dashboard"]["B6"].value)
    print("  Dashboard B7:", wb2["Dashboard"]["B7"].value)
    print("  Charts:", sum(len(wb2[s]._charts) for s in wb2.sheetnames))
    print("  Tables:", sum(len(wb2[s].tables) for s in wb2.sheetnames))


if __name__ == "__main__":
    main()
