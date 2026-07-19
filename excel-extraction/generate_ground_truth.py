"""Generate ground truth for the complex test workbook.

Captures static cell values, formulas, computed formula values, merges,
tables, charts, hyperlinks, comments, and defined names from
`test_data/complex_workbook.xlsx`.

Formula values are computed using the `formulas` library so libraries that
only read cached values can be benchmarked against ground truth.
"""

from __future__ import annotations

import json
import pathlib
import re

import formulas
import openpyxl
import numpy as np


HERE = pathlib.Path(__file__).parent
WORKBOOK = HERE / "test_data" / "complex_workbook.xlsx"
GT_DIR = HERE / "ground_truth"
GT_DIR.mkdir(parents=True, exist_ok=True)

SHEET_KEY_RE = re.compile(r"^'\[([^\]]+)\]([^!]+)!([^!]+)$")


def _norm_date(v):
    import datetime as _dt
    if isinstance(v, _dt.datetime):
        return v.date().isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    return v


def _norm(v):
    if v is None:
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (np.ndarray,)):
        return [_norm(x) for x in v.tolist()]
    if isinstance(v, list):
        return [_norm(x) for x in v]
    if isinstance(v, tuple):
        return [_norm(x) for x in v]
    if isinstance(v, dict):
        return {k: _norm(x) for k, x in v.items()}
    if hasattr(v, "isoformat"):
        return _norm_date(v)
    return v


def compute_values() -> dict[str, dict[str, object]]:
    xl = formulas.ExcelModel().loads(str(WORKBOOK)).finish()
    sol = xl.calculate()
    out: dict[str, dict[str, object]] = {}
    for k, node in sol.items():
        m = SHEET_KEY_RE.match(k)
        if not m:
            continue
        sheet = m.group(2).rstrip("'")
        coord = m.group(3)
        try:
            val = node.value
            if hasattr(val, "tolist"):
                val = val.tolist()
            if isinstance(val, list) and len(val) == 1 and isinstance(val[0], list) and len(val[0]) == 1:
                val = val[0][0]
        except Exception:
            val = None
        out.setdefault(sheet, {})[coord] = val
    return out


def chart_info(ch):
    try:
        series_count = len(ch.series) if hasattr(ch, "series") else 0
    except Exception:
        series_count = 0
    info = {"type": type(ch).__name__, "title": None, "series_count": series_count}
    try:
        if ch.title and hasattr(ch.title, "tx") and ch.title.tx and ch.title.tx.rich:
            paras = ch.title.tx.rich.paragraphs
            if paras and paras[0].r:
                info["title"] = paras[0].r[0].t
    except Exception:
        pass
    return info


def build_ground_truth() -> dict:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    computed = compute_values()
    gt: dict = {"sheets": {}}

    for sn in wb.sheetnames:
        ws = wb[sn]
        sheet_gt = {
            "dimensions": {"max_row": ws.max_row, "max_column": ws.max_column},
            "frozen_panes": ws.freeze_panes,
            "merged_cells": sorted(str(r) for r in ws.merged_cells.ranges),
            "tables": [],
            "charts": [],
            "hyperlinks": [],
            "comments": [],
            "column_widths": {},
            "cells": {},
        }

        for tbl_name, tbl_ref in ws.tables.items():
            sheet_gt["tables"].append({"name": tbl_name, "ref": tbl_ref})

        for ch in ws._charts:
            sheet_gt["charts"].append(chart_info(ch))

        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None:
                    sheet_gt["hyperlinks"].append({
                        "cell": cell.coordinate,
                        "target": getattr(cell.hyperlink, "target", None) or getattr(cell.hyperlink, "location", None),
                        "display": cell.value,
                    })
                if cell.comment is not None:
                    sheet_gt["comments"].append({
                        "cell": cell.coordinate,
                        "author": cell.comment.author,
                        "text": cell.comment.text,
                    })

        for col_letter, dim in ws.column_dimensions.items():
            if dim.width:
                sheet_gt["column_widths"][col_letter] = dim.width

        sheet_comp = computed.get(sn.upper(), {})
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    sheet_gt["cells"][cell.coordinate] = {
                        "value": _norm(sheet_comp.get(cell.coordinate)),
                        "formula": cell.value,
                        "data_type": cell.data_type,
                    }
                else:
                    sheet_gt["cells"][cell.coordinate] = {
                        "value": _norm(cell.value),
                        "formula": None,
                        "data_type": cell.data_type,
                    }

        gt["sheets"][sn] = sheet_gt

    gt["defined_names"] = {n: wb.defined_names[n].attr_text for n in wb.defined_names}
    return gt


def main():
    gt = build_ground_truth()
    out_path = GT_DIR / "ground_truth.json"
    out_path.write_text(json.dumps(gt, indent=2, default=str))
    print(f"Wrote {out_path}")
    for sn, data in gt["sheets"].items():
        ncells = len(data["cells"])
        nformulas = sum(1 for c in data["cells"].values() if c["formula"])
        nwith_val = sum(1 for c in data["cells"].values() if c["formula"] and c["value"] is not None)
        print(f"  {sn}: {ncells} cells, {nformulas} formulas ({nwith_val} with computed values), "
              f"{len(data['charts'])} charts, {len(data['tables'])} tables, {len(data['merged_cells'])} merges, "
              f"{len(data['hyperlinks'])} hyperlinks, {len(data['comments'])} comments")
    print(f"Defined names: {list(gt['defined_names'])}")


if __name__ == "__main__":
    main()
